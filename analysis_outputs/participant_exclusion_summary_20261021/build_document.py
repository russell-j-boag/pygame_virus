from pathlib import Path
import csv
import sys
import hashlib
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/rjb779/Documents/pygame_virus_aid_onset')
OUT = Path(__file__).parent
QC = ROOT / 'analysis_outputs/semester2_2026_qc'
sys.path.insert(0, '/Users/rjb779/.codex/plugins/cache/openai-primary-runtime/documents/26.623.12021/skills/documents/scripts')
from table_geometry import apply_table_geometry

# Preset: standard_business_brief; header: memo_masthead without decorative rule.
# Named overrides: TableText 10 pt; FigureCaption 10 pt; landscape figure pages
# use 0.5-inch vertical and 0.75-inch horizontal margins for readable source plots.
doc = Document()
def style(name, size, before=0, after=6, color='222222', bold=False, line=1.1):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = bold
    s.font.color.rgb = RGBColor.from_string(color)
    f=s.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line
    return s
style('Normal',11)
style('Title',23,after=4,color='172B3A',bold=True)
style('Subtitle',12,after=12,color='555555')
for n,size,before,after,color in [(1,16,16,8,'2E74B5'),(2,13,12,6,'2E74B5'),(3,12,8,4,'1F4D78')]:
    style(f'Heading {n}',size,before,after,color,True).paragraph_format.keep_with_next=True
style('TableText',10,after=3)
style('SourceNote',9,before=4,after=4,color='555555')
style('FigureCaption',10,after=5)
style('PlotImage',10,after=0,line=1)
style('Header',9,after=0,color='666666')
style('Footer',9,after=0,color='666666')
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.492)
s.header.paragraphs[0].text='Aid-onset study | Participant exclusion review'
p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('Page ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
doc.core_properties.title='Aid-onset study: participant exclusion review'
doc.core_properties.subject='Recorded candidate status and individual diagnostic plots; prepared 21 October 2026'

import json

def read_csv(path):
    with path.open(newline='') as f:
        return list(csv.DictReader(f))

summary = json.loads((QC/'exclusion_summary.json').read_text())
metrics = read_csv(QC/'scope_metrics.csv')
flags = read_csv(QC/'flagged_checks.csv')
candidates = read_csv(QC/'exclusion_candidates.csv')
historical = read_csv(ROOT/'analysis_outputs/semester2_2026_hypotheses/analysis_exclusions.csv')
ids = [int(r['participant_id']) for r in candidates]
assert ids == [10,13,20,22,27] == summary['participant_ids']
assert summary['exclusions_applied'] is False and summary['participant_n'] == 60
assert historical[0]['excluded_run_timestamp'] == '20260903_120309'
assert historical[0]['replacement_run_timestamp'] == '20260924_110831'
trigger_flags = [r for r in flags if r['use_for_exclusion'] == 'True']
assert {r['decision'] for r in trigger_flags} == {'1'}
names = {'MANUAL':'Manual', 'STIMFIRST':'Stimulus-first', 'aided_overall':'Pooled aided', 'experimental_overall':'All main trials'}
scopes = ['MANUAL','STIMFIRST','aided_overall','experimental_overall']

doc.add_paragraph('Participant exclusion review','Title')
doc.add_paragraph('Aid-onset study | Recorded candidate status | 21 October 2026','Subtitle')
doc.add_paragraph('Five of 60 participants (8.3%) met the agreed screening criteria. All five remain included in the current analysis; the saved audit records candidates, not applied exclusions. All candidate flags arose in Decision 1. Excluding all five would leave 55 participants.')

rows = [['Participant','Candidate reason','Supporting diagnostic evidence (Decision 1)']]
for pid in ids:
    pf = [r for r in trigger_flags if int(r['participant_id']) == pid]
    near = any(r['category']=='near_chance_or_worse' for r in pf)
    reason = 'Anticipatory responding' + (' and near-chance performance' if near else '')
    evidence = []
    for scope in scopes:
        sf = [r for r in pf if r['scope']==scope]
        if not sf: continue
        m = next(r for r in metrics if int(r['participant_id'])==pid and r['scope']==scope and r['decision']=='1')
        text = f"{names[scope]}: {m['n_fast']}/{m['n_trials']} fast ({float(m['fast_fraction'])*100:.1f}%); fast accuracy {float(m['fast_accuracy'])*100:.1f}%."
        if any(r['category']=='near_chance_or_worse' for r in sf):
            text += f" Overall accuracy {float(m['accuracy'])*100:.1f}%."
        evidence.append(text)
    rows.append([f'{pid:03d}',reason,'\n'.join(evidence)])

t=doc.add_table(rows=len(rows),cols=3); t.style='Table Grid'
for i,values in enumerate(rows):
    trpr=t.rows[i]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    for j,value in enumerate(values):
        cell=t.cell(i,j); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.style='TableText'; p.add_run(value)
        if j==0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if i==0:
            p.runs[0].bold=True
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F2F4F7'); cell._tc.get_or_add_tcPr().append(sh)
    if i==0: trpr.append(OxmlElement('w:tblHeader'))
apply_table_geometry(t,[1152,2376,5832])
doc.add_paragraph('Fast = recorded RT <200 ms. Fast percentages use all trials in the stated scope; fast accuracy uses only fast responses. Pooled scopes overlap condition scopes and do not represent additional participants. Source: saved QC scope metrics and flagged checks.','SourceNote')
doc.add_paragraph('The earlier participant-059 run was excluded separately for chance performance; its replacement is retained (see next page).','SourceNote')

doc.add_page_break()
doc.add_paragraph('Recorded run exclusion: participant 059','Heading 1')
doc.add_paragraph('The analysis exclusion record identifies the earlier run 20260903_120309 as excluded for chance performance. The replacement run 20260924_110831 is included in the canonical 60-person cohort. This is a run-specific exclusion, not an exclusion of participant 059 from the current analysis.')
m59=next(r for r in metrics if r['participant_id']=='59' and r['scope']=='MANUAL' and r['decision']=='2')
assert m59['run_timestamp']=='20260924_110831'
doc.add_paragraph(f"The replacement is not an exclusion candidate. Its Manual Decision 2 accuracy was {m59['n_correct']}/{m59['n_trials']} ({float(m59['accuracy'])*100:.1f}%). Its practice and Manual target-deviation flags are diagnostic only. The historical run is outside the current QC bundle and the five-candidate total.")

doc.add_paragraph('Repeated IDs in the raw archive','Heading 2')
doc.add_paragraph('The current analysis dataset contains 60 distinct participant IDs and 60 selected runs: one run per ID, with no duplicated participant/run/block/trial records. The raw Semester 2 archive contains 63 distinct ID/run pairs across those 60 IDs. Three IDs have more than one recorded run:')
for label, text in [
    ('014', 'Earlier run 20260727_110459 is incomplete: practice and two main-block files are present, but there is no combined ALL export. Later complete run 20260730_085751 is selected.'),
    ('053', 'Earlier run 20260827_125739 contains only a practice file. Later complete run 20260827_131135 is selected.'),
    ('059', 'Earlier complete run 20260903_120309 was excluded for chance performance. Later complete run 20260924_110831 is selected, as described above.')
]:
    p=doc.add_paragraph(); p.add_run(f'Participant {label}. ').bold=True; p.add_run(text)
doc.add_paragraph('These are repeated IDs across raw runs, not duplicate IDs among the selected analysis runs. The files alone do not establish whether different people used the same ID. Earlier incomplete runs for 014 and 053 are not additional behavioural exclusion candidates.')
doc.add_paragraph('Identity check sources: data/data_virus_all.csv; data/collation_manifest.csv; CSV filenames and row metadata in output/semester2_2026_data/. The collation script selects the latest available complete export set for each ID and stops if that selected run is missing an export family.','SourceNote')

doc.add_page_break()
doc.add_paragraph('Remaining participant IDs to collect', 'Heading 2')
collected_ids = {int(r['participant_id']) for r in read_csv(ROOT/'data/collation_manifest.csv')}
assert collected_ids == set(range(1, 61))
target_n = 60
new_ids = sorted(set(range(1, target_n + 1)) - collected_ids)
p = doc.add_paragraph()
p.add_run('Re-collect: ').bold = True
p.add_run(', '.join(f'{pid:03d}' for pid in ids) + ' (5 participants, if the candidate exclusions are confirmed).')
p = doc.add_paragraph()
p.add_run('New participants: ').bold = True
p.add_run(', '.join(f'{pid:03d}' for pid in new_ids) if new_ids else 'None. All IDs 001-060 have been collected.')
doc.add_paragraph('Total remaining: 5 re-collections if all five candidates are excluded and replaced, to retain the target of 60 participants. The five candidates remain included in the current analysis. Participant 059 already has a retained replacement and does not need re-collection.')
doc.add_paragraph('Collection status checked against data/collation_manifest.csv on 21 October 2026. Target: 60 retained participants.', 'SourceNote')
doc.add_paragraph('Review criteria','Heading 2')
doc.add_paragraph('Anticipatory responding required ≥20% of all scope trials with RT <200 ms, ≥20 fast responses, and ≤55% accuracy on those fast responses, jointly within the same decision and scope. Near-chance performance meant ≤55% overall accuracy, counting timeouts as incorrect. These retrospective descriptive screens were agreed on 17 October 2026; they were not preregistered exclusions or evidence of equivalence to chance.')
doc.add_paragraph('Decision 1 and Decision 2 were screened separately for each 260-trial main condition, the 520 pooled aided trials and all 780 main trials. Timeouts never count as fast; exactly 200 ms is not fast. Short RTs require interpretation: Aid-first Decision 1 can be prepared from preceding advice, and Decision 2 can be prepared before its response screen.')
doc.add_paragraph('Practice-target deviations (11 participants) and Manual Decision 2 target deviations (15 participants) were diagnostic only, using an absolute deviation >10 percentage points from the recorded 75% target. After overlap with each other and the candidates, 18 participants had diagnostic-only flags. All assessed scopes were assessable.')

doc.add_paragraph('Reading the diagnostic plots','Heading 2')
doc.add_paragraph('The following pages reproduce the saved individual plots. Each shows both decisions across 60 practice trials and 780 main trials in administered order. Headlines pool main trials; the table lists the scopes that triggered screening. The “candidate” label applies to the participant, including above the Decision 2 panels; no Decision 2 scope triggered a candidate flag.')
doc.add_paragraph('Practice bins contain 10 trials; main bins contain 50 trials, with a final 10-trial bin. Grey bands show condition-aligned peer medians and interquartile ranges. Binned patterns and peer comparisons are descriptive, not separate exclusion rules.')
doc.add_paragraph('Sources: analysis_outputs/semester2_2026_qc/ (exclusion_summary.json, exclusion_candidates.csv, scope_metrics.csv, flagged_checks.csv and individuals/participant_###.png); analysis_outputs/semester2_2026_hypotheses/analysis_exclusions.csv; README.md. Prepared from saved outputs on 21 October 2026. Source data, screening decisions and analyses were not changed.','SourceNote')

caps={10:'Manual, Stimulus-first and all main trials met the anticipatory criterion (Decision 1).',13:'Stimulus-first met the anticipatory criterion (Decision 1).',20:'Manual met the anticipatory criterion (Decision 1).',22:'Stimulus-first met both criteria; Manual and both pooled scopes also met the anticipatory criterion (Decision 1).',27:'Manual met both the anticipatory and near-chance criteria (Decision 1).'}
for idx,pid in enumerate(ids):
    if idx==0:
        s=doc.add_section(WD_SECTION_START.NEW_PAGE)
        s.orientation=WD_ORIENT.LANDSCAPE; s.page_width=Inches(11); s.page_height=Inches(8.5)
        s.top_margin=s.bottom_margin=Inches(.5); s.left_margin=s.right_margin=Inches(.75)
        s.header_distance=s.footer_distance=Inches(.25)
    else: doc.add_page_break()
    p=doc.add_paragraph(style='FigureCaption'); p.paragraph_format.keep_with_next=True
    p.add_run(f'Participant {pid:03d}. ').bold=True; p.add_run(caps[pid]+' Status: candidate; retained.')
    p=doc.add_paragraph(style='PlotImage'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=p.add_run().add_picture(str(QC/f'individuals/participant_{pid:03d}.png'),width=Inches(9.2))
    shape._inline.docPr.set('descr',f'Participant {pid:03d}: separate Decision 1 and Decision 2 RTs, fast-response rates, all-trial accuracy and fast-only accuracy across practice and main conditions.')

target=OUT/'aid_onset_reviewed_exclusions.docx'
doc.save(target)
with ZipFile(target) as z:
    media=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
    assert len(media)==5
    for pid in ids: assert (QC/f'individuals/participant_{pid:03d}.png').read_bytes() in media
assert len(Document(target).tables)==1
sources=[QC/n for n in ['exclusion_summary.json','exclusion_candidates.csv','scope_metrics.csv','flagged_checks.csv']]
sources += [ROOT/'README.md',ROOT/'analysis_outputs/semester2_2026_hypotheses/analysis_exclusions.csv']
sources += [QC/f'individuals/participant_{pid:03d}.png' for pid in ids]
(OUT/'source_checksums.json').write_text(json.dumps({str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},indent=2)+'\n')
print(target)
print('Verified 5 candidate IDs; retained status; run-specific p059 exclusion; 5 byte-identical source plots.')
